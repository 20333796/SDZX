import hashlib
import json
import re
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.models import KnowledgeChunkModel, KnowledgeDocumentModel
from app.schemas import DocumentStatus
from app.services.embeddings import EmbeddingError, EmbeddingProvider
from app.storage import ObjectStore

MAX_PARSE_ATTEMPTS = 3
CHUNK_TARGET_SIZE = 900


def extract_document_text(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return payload.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("Text document encoding is not supported")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(payload)).pages)
    if suffix == ".docx":
        from docx import Document

        return "\n".join(paragraph.text for paragraph in Document(BytesIO(payload)).paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation

        return "\n".join(
            shape.text
            for slide in Presentation(BytesIO(payload)).slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )
    raise ValueError("Unsupported document format")


def chunk_text(text: str, target_size: int = CHUNK_TARGET_SIZE) -> list[str]:
    paragraphs = [re.sub(r"\s+", " ", item).strip() for item in re.split(r"\n{2,}", text) if item.strip()]
    chunks: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        if len(paragraph) > target_size:
            if buffer:
                chunks.append(buffer)
                buffer = ""
            chunks.extend(paragraph[index:index + target_size] for index in range(0, len(paragraph), target_size))
            continue
        if buffer and len(buffer) + len(paragraph) + 1 > target_size:
            chunks.append(buffer)
            buffer = paragraph
        else:
            buffer = f"{buffer} {paragraph}".strip()
    if buffer:
        chunks.append(buffer)
    return chunks


def parse_document(session: Session, store: ObjectStore, document: KnowledgeDocumentModel) -> int:
    if document.status != DocumentStatus.uploaded.value:
        return 0
    document.parsing_attempts += 1
    try:
        text = extract_document_text(document.original_filename, store.get(document.object_key))
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("Document contains no extractable text")
        session.execute(delete(KnowledgeChunkModel).where(KnowledgeChunkModel.document_id == document.id))
        for ordinal, content in enumerate(chunks):
            session.add(
                KnowledgeChunkModel(
                    id=str(uuid4()),
                    document_id=document.id,
                    ordinal=ordinal,
                    content=content,
                    content_hash=hashlib.sha256(content.encode("utf-8")).hexdigest(),
                    source_locator=f"chunk:{ordinal + 1}",
                )
            )
        document.status = DocumentStatus.parsed.value
        document.processing_error = None
        document.parsed_at = datetime.now(UTC)
        session.commit()
        return len(chunks)
    except Exception as error:
        document.processing_error = str(error)[:2000]
        if document.parsing_attempts >= MAX_PARSE_ATTEMPTS:
            document.status = DocumentStatus.rejected.value
            document.reviewed_by = "ingestion-worker"
            document.reviewed_at = datetime.now(UTC)
        session.commit()
        return 0


def process_pending_documents(session: Session, store: ObjectStore, batch_size: int = 10) -> int:
    documents = session.scalars(
        select(KnowledgeDocumentModel)
        .where(KnowledgeDocumentModel.status == DocumentStatus.uploaded.value)
        .order_by(KnowledgeDocumentModel.created_at)
        .limit(batch_size)
    ).all()
    return sum(parse_document(session, store, document) for document in documents)


def index_reviewed_documents(session: Session, provider: EmbeddingProvider, batch_size: int = 10) -> int:
    documents = session.scalars(
        select(KnowledgeDocumentModel)
        .where(KnowledgeDocumentModel.status == DocumentStatus.reviewed.value)
        .order_by(KnowledgeDocumentModel.reviewed_at, KnowledgeDocumentModel.created_at)
        .limit(batch_size)
    ).all()
    indexed = 0
    for document in documents:
        chunks = session.scalars(
            select(KnowledgeChunkModel)
            .where(KnowledgeChunkModel.document_id == document.id)
            .order_by(KnowledgeChunkModel.ordinal)
        ).all()
        if not chunks:
            document.processing_error = "Document has no parsed knowledge chunks"
            session.commit()
            continue
        try:
            vectors = provider.embed([chunk.content for chunk in chunks])
            if len(vectors) != len(chunks):
                raise EmbeddingError("Embedding service returned an incomplete batch")
            for chunk, vector in zip(chunks, vectors, strict=True):
                chunk.embedding = json.dumps(vector, separators=(",", ":"))
                chunk.embedding_model = provider.model_name
            document.status = DocumentStatus.indexed.value
            document.processing_error = None
            session.commit()
            indexed += len(chunks)
        except EmbeddingError as error:
            document.processing_error = str(error)[:2000]
            session.commit()
    return indexed
